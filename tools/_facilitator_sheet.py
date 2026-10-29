import json, docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAMS = json.load(open("_teams.json"))
HEAD, BODY = "Plus Jakarta Sans", "Inter"
INK, MUTE, GREEN = RGBColor(0x11,0x11,0x11), RGBColor(0x6b,0x72,0x80), RGBColor(0x2E,0x7D,0x46)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.5)
st = doc.styles["Normal"]; st.font.name = BODY; st.font.size = Pt(10); st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.12
st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

def shade(cell, hexcolor):
    el = OxmlElement('w:shd'); el.set(qn('w:val'),'clear'); el.set(qn('w:fill'), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)

def h(text, size=20, color=INK, before=10, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.font.name = HEAD; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = color
    r._element.rPr.rFonts.set(qn('w:eastAsia'), HEAD)
    return p

def para(text, size=10, color=INK, italic=False, bold=False, after=4, space_before=0):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(space_before)
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color
    r.font.italic = italic; r.font.bold = bold
    return p

def table(headers, rows, widths, header_fill="1C5230"):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # LibreOffice ignores cell widths unless autofit is off and the grid columns carry them too.
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Inches(w)
    for i, (c, w) in enumerate(zip(t.rows[0].cells, widths)):
        c.width = Inches(w); shade(c, header_fill)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        r = c.paragraphs[0].add_run(headers[i]); r.font.bold = True; r.font.size = Pt(9)
        r.font.name = HEAD; r.font.color.rgb = RGBColor(0xff,0xff,0xff)
    for row in rows:
        cells = t.add_row().cells
        for i, (c, val) in enumerate(zip(cells, row)):
            c.width = Inches(widths[i])
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
            txt, bold, fill = (val if isinstance(val, tuple) else (val, False, None))
            if fill: shade(c, fill)
            r = c.paragraphs[0].add_run(txt); r.font.size = Pt(9); r.font.bold = bold
            if fill: r.font.color.rgb = RGBColor(0xff,0xff,0xff)
    return t

# ── header ────────────────────────────────────────────────────────────────
h("Field Guide Challenge: Facilitator Sheet", 22, GREEN, before=0, after=2)
para("Liberty Utilities Central · Tree Line USA, October 2026 · keep this sheet face down",
     9, MUTE, italic=True, after=10)

# ── check-in ──────────────────────────────────────────────────────────────
h("1. Check-in: finish codes", 14)
para("Every team finishes on its own: their last screen shows a 4-digit finish code. Check it "
     "against the team name, then write the next number in the Finished column. Use as many"
     " QR cards as you have teams, in any order.", 9.5, after=6)
ck = table(["Team", "Finish code", "Finished #", "Stop order"],
           [[(t["name"], True, None), (t["code"], True, None), "", t["ordernames"]] for t in TEAMS],
           [1.9, 1.1, 1.1, 3.4])
from docx.enum.table import WD_ROW_HEIGHT_RULE
for row in ck.rows[1:]:       # tall enough to write a number in by hand
    row.height = Inches(0.42); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
para("A code that doesn't match the team means they read it off another team's screen.",
     9, MUTE, space_before=4, after=8)

# ── per-team answers ──────────────────────────────────────────────────────
doc.add_page_break()          # check-in is the page Lindsey keeps on top; answers go underneath
h("2. Each team's answers", 14, before=0)
para("Every team does all five stops, each in its own order, so the room starts spread out and the "
     "plant wall never gets everyone at once. Three stops differ by team: each team hunts its "
     "plant on the wall (two pairs share one) and gets one of two field clips and one of three field notes, so "
     "neighboring teams are not chasing the same answer. Search (HOLLY) and Field Bites (FOXTAILS) "
     "are the same for everyone.", 9.5, after=6)
table(["Team", "Plant ID: plant (sheet), height", "Field Clips: answer", "Field Notes: answer"],
      [[(t["name"], True, None),
        "%s, sheet %d: %s" % (t["plantname"], t["sheet"], t["height"] if t["checked"] else "any number"),
        "%s: %s" % (t["cliplabel"], t["clipanswer"]),
        "%s: %s" % (t["notelabel"], t["noteanswer"])] for t in TEAMS],
      [1.55, 2.2, 1.95, 1.8])
para("", 6, after=4)

# ── the five stops ────────────────────────────────────────────────────────
h("3. The five stops and how to unstick a team", 14)
para("Teams count stops in their own order, so \u201cwe\u2019re stuck on stop 2\u201d means a "
     "different stop for each team. Ask which section is named at the top of their screen. If a "
     "team has been stuck a few minutes, give the nudge, not the answer.", 9.5, after=6)

STOPS = [
 ("Search",
  "Homeowner wants a bird-friendly shrub off to the side of the line. Filters, in the app's order: slow growing, wire zone no / border zone yes, host plant and bird food, 15 ft or less. One plant fits.",
  "HOLLY",
  "\u201cDid you set all four filters?\u201d Leaving one off brings up more than one plant. The zone filter is one choice: wire zone no, border zone yes."),
 ("Plant ID",
  "Find their own plant on the wall (named on their screen) by taking pictures with Plant ID, then read its max height off its plant profile.",
  "That team's height (section 2)",
  "Ask which plant part they chose (leaf, flower, fruit or bark); most bad matches are the wrong part. They type the feet (12 ft 0 in reads as 12). Use the height on the plant profile, not a number from a field bite. If they truly cannot find it, tell them the sheet number."),
 ("Field Clips",
  "Rose clip: the invasive rose's little leaflet at the base of each leaf is covered in what?\nVine clip: Virginia creeper's tendrils end in sticky what?",
  "Rose: EYELASHES\nVine: PADS",
  "Rose: \u201cWatch the part about the base of the leaf, about two-thirds of the way in.\u201d\nVine: \u201cIt comes right after the part about poison ivy\u2019s hairy roots.\u201d"),
 ("Field Notes",
  "Johnsongrass: how far can its rhizomes travel from the mother plant, in feet?\nElderberry Bloom: the note names bees, flies and what other insect?\nAutumn Olive: wanted dead or what?",
  "Johnsongrass: 200\nElderberry: BEETLES\nAutumn Olive: CONTROLLED",
  "Johnsongrass: \u201cLook in the Reproduction box.\u201d\nElderberry: \u201cRead the paragraph under the name.\u201d\nAutumn Olive: \u201cIt's the headline.\u201d"),
 ("Field Bites",
  "Work through Black Locust vs False Indigo Bush. False indigo bush flowers look like little purple what?",
  "FOXTAILS",
  "\u201cOn the Flowers card, read the line under the heading, not the heading.\u201d"),
]
table(["Section", "What they have to do", "Answer", "If they\u2019re stuck"],
      [[(s[0], True, None), s[1], (s[2], True, None), s[3]] for s in STOPS],
      [1.0, 2.55, 1.5, 2.45])
para("Answers ignore capitals, spaces and punctuation.", 9, MUTE, space_before=4, after=8)

doc.add_page_break()

# ── plant wall ────────────────────────────────────────────────────────────
h("4. Plant wall key", 14, before=0)
WALL = json.load(open("_wallkey.json"))
para(f"{len(WALL)} sheets, one plant each, numbered but deliberately NOT labeled. Sheets 4 and 10 "
     "came off the wall on October 29 (they did not scan reliably), so two sheets are shared by two "
     "teams. This is the only key.", 9.5, after=4)
para("The photos are a MIX of flower, fruit, leaf and bark, because Plant ID asks which one you are "
     "photographing. Choosing the right plant part is the skill this stop teaches. The \u201cPhoto is "
     "a\u201d column is the plant part they should choose.", 9.5, after=6)
owner = {}
for t in TEAMS:
    owner[t["sheet"]] = owner[t["sheet"]] + " and " + t["name"] if t["sheet"] in owner else t["name"]
table(["Sheet", "What it is", "Photo is a", "Whose plant", "Notes"],
      [[(str(w["n"]), True, None), (w["name"], True, None), (w["organ"], True, None),
        owner.get(w["n"], ""), w.get("note", "")] for w in WALL],
      [0.65, 1.5, 0.85, 1.6, 2.9])
para("", 6, after=6)

# ── if it goes wrong ──────────────────────────────────────────────────────
h("5. If something goes wrong", 14)
for problem, fix in [
  ("A team's QR won't scan", "Type the address on the bottom of their card straight into the browser."),
  ("They closed the page / phone died", "Reopening the same link puts them back where they were; progress is saved on that phone. If they switch phones they start over, so keep one phone driving."),
  ("The app won't identify a printed plant", "Nine times out of ten they chose the wrong plant part (leaf, flower, fruit or bark). Check that first. If it still will not match, give them their plant's height from section 2 and move them on."),
  ("They're not logged in to Let Grow", "That is the point of the game. Get them logged in before they start."),
  ("A team is on another team's page", "Their team name is in the bar at the top of every screen. If it isn't theirs, have them scan their own card or type the address printed on it. If they already finished, write the finish under their real team."),
  ("A team is miles behind", "Give them the answer to a stop and move them on so they still finish."),
  ("Two teams arrive at once", "Check both codes and write them down in the order they reached you."),
]:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(problem + " — "); r.font.bold = True; r.font.size = Pt(9.5)
    r2 = p.add_run(fix); r2.font.size = Pt(9.5)

para("Everything the teams need is inside the Let Grow app. If the app is down, the game is down, "
     "so check it the morning of.", 9, MUTE, space_before=8)

doc.save("Facilitator Sheet.docx")
print("built Facilitator Sheet.docx")
